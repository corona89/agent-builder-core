from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

def create_strategy_ppt():
    prs = Presentation()
    # 16:9 widescreen layout set
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide_layout = prs.slide_layouts[5] # Blank layout with title
    slide = prs.slides.add_slide(slide_layout)
    
    # Background (Dark Zinc/Blue tone like the webpage)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(9, 9, 11) # bg-zinc-950
    
    # Title
    title_shape = slide.shapes.title
    title_shape.text = "2026 AI:ON-U 고도화 전략 및 실행 계획"
    title_tf = title_shape.text_frame
    title_tf.paragraphs[0].font.size = Pt(36)
    title_tf.paragraphs[0].font.bold = True
    title_tf.paragraphs[0].font.color.rgb = RGBColor(59, 130, 246) # Blue-500
    title_tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    
    # Helper to add section box
    def add_section(left_inch, top_inch, width_inch, height_inch, title, items, color_rgb):
        # Header line
        shape = slide.shapes.add_textbox(Inches(left_inch), Inches(top_inch), Inches(width_inch), Inches(0.4))
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = color_rgb
        
        # Content box
        content_shape = slide.shapes.add_textbox(Inches(left_inch), Inches(top_inch + 0.5), Inches(width_inch), Inches(height_inch))
        tf = content_shape.text_frame
        tf.word_wrap = True
        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(212, 212, 216) # zinc-300
            p.space_after = Pt(5)

    # Left Column: 전략 수립 체계
    add_section(0.5, 1.5, 6, 2.5, "1. 고도화 전략 수립 체계", [
        "기술 경쟁력: Graph RAG 도입, 멀티모달(오디오/영상) 지원, A2A 표준 연동",
        "사업 경쟁력: 대규모 트래픽 안정성, 모델 관리 도구 구현, 가드레일 서비스",
        "AI Guild: 전문가 그룹 기반 신기술 PoC 및 현장 피드백 반영 체계 운영"
    ], RGBColor(59, 130, 246))

    # Right Column: 실행 계획
    add_section(6.8, 1.5, 6, 2.5, "2. 고도화 실행 계획", [
        "보안성 강화: 채팅창 로그인(ABCLab, 이메일 OTP), 배포 승인 절차 최적화",
        "지능화: 플래닝 에이전트 기반 워크플로우 자동화, Human-in-the-loop 도입",
        "데이터: 문서 메타데이터 처리 고도화, 온톨로지 기반 지식 관리 통합"
    ], RGBColor(147, 51, 234)) # Purple-600

    # Bottom: Goals
    add_section(0.5, 4.5, 12.3, 2, "3. 주요 목표 (Goals)", [
        "사업: 가드레일(자체/글로벌) & 에이전트 평가 자동화 도구 상용 수준 고도화",
        "기술: 에이전트 표준(MCP, A2A) 전면 지원 및 멀티모달 지식 관리 최적화",
        "사용성: 플래닝 에이전트 기반 커스텀 에이전트 제작 환경 자동화 및 대중화"
    ], RGBColor(16, 185, 129)) # Emerald-500

    # Save
    output_path = "strontium-webpage/backend/uploads/AION_U_Strategy_1Page.pptx"
    prs.save(output_path)
    print(output_path)

if __name__ == "__main__":
    create_strategy_ppt()
